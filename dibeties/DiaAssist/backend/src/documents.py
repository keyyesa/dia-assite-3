"""Text extraction for uploaded files (knowledge-base documents and lab reports)."""
import os
import re

import config

try:  # PyMuPDF >= 1.24 exposes `pymupdf`; older builds only expose `fitz`
    import pymupdf as _pymupdf
except ImportError:  # pragma: no cover
    import fitz as _pymupdf


IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg")
REPORT_EXTENSIONS = tuple(config.SUPPORTED_EXTENSIONS) + IMAGE_EXTENSIONS  # lab reports may also be photos


def ocr_available() -> bool:
    """True when Tesseract OCR (pytesseract + the Tesseract program) can be used."""
    try:
        import pytesseract
        _configure_tesseract(pytesseract)
        pytesseract.get_tesseract_version()
        return True
    except Exception:  # noqa: BLE001 - missing package or program
        return False


def _configure_tesseract(pytesseract) -> None:
    cmd = os.getenv("TESSERACT_CMD", "").strip()
    if cmd:
        pytesseract.pytesseract.tesseract_cmd = cmd


_OCR_HELP = ("This file has no text layer, so OCR is needed. Install the Tesseract program "
             "(Windows: github.com/UB-Mannheim/tesseract/wiki), run `pip install pytesseract pillow`, "
             "and set TESSERACT_CMD in .env if it is not on PATH.")


def _ocr_png_bytes(png: bytes) -> str:
    import io
    import pytesseract
    from PIL import Image
    _configure_tesseract(pytesseract)
    return pytesseract.image_to_string(Image.open(io.BytesIO(png)))


def _ocr_image(data: bytes) -> str:
    if not ocr_available():
        raise ValueError(_OCR_HELP)
    return _ocr_png_bytes(data)


def _ocr_pdf(data: bytes) -> str:
    if not ocr_available():
        raise ValueError(_OCR_HELP)
    doc = _pymupdf.open(stream=data, filetype="pdf")
    try:
        return "\n".join(_ocr_png_bytes(page.get_pixmap(dpi=200).tobytes("png")) for page in doc)
    finally:
        doc.close()


def is_supported(filename: str) -> bool:
    return os.path.splitext(filename.lower())[1] in config.SUPPORTED_EXTENSIONS


def safe_filename(name: str) -> str:
    """Strip any path and unsafe characters from an uploaded file name."""
    base = os.path.basename((name or "").replace("\\", "/")).strip()
    base = re.sub(r"[^A-Za-z0-9._ ()\-]", "_", base).lstrip(".").strip()
    stem, ext = os.path.splitext(base)
    return (stem[:120] + ext.lower()) if stem else ""


def _pdf_text(data: bytes, page_markers: bool) -> str:
    doc = _pymupdf.open(stream=data, filetype="pdf")
    try:
        if getattr(doc, "needs_pass", False):
            raise ValueError("This PDF is password-protected.")
        pages = []
        for i, page in enumerate(doc):
            text = page.get_text()
            pages.append(f"[Page {i + 1}]\n{text}" if page_markers else text)
        return "\n".join(pages)
    finally:
        doc.close()


def _docx_text(data: bytes) -> str:
    try:
        import io
        import docx  # python-docx
    except ImportError as exc:
        raise ValueError("Reading .docx files needs python-docx: pip install python-docx") from exc
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes, page_markers: bool = True, ocr: bool = False) -> str:
    """Return the plain text of a PDF / TXT / MD / DOCX file. Raises ValueError with a
    human-readable message when the file cannot be read.
    ocr=True (used for lab reports) also reads photos (PNG/JPG) and scanned PDFs through Tesseract."""
    ext = os.path.splitext(filename.lower())[1]
    try:
        if ocr and ext in IMAGE_EXTENSIONS:
            return _ocr_image(data)
        if ext == ".pdf":
            text = _pdf_text(data, page_markers)
            if ocr and not re.sub(r"\[Page \d+\]|\s", "", text):
                return _ocr_pdf(data)  # no text layer: a scanned report
            return text
        if ext in (".txt", ".md"):
            return data.decode("utf-8", errors="ignore")
        if ext == ".docx":
            return _docx_text(data)
    except ValueError:
        raise
    except Exception as exc:  # noqa: BLE001 - corrupt files raise many different errors
        raise ValueError(f"Could not read '{filename}': {exc}") from exc
    raise ValueError(
        f"Unsupported file type '{ext or filename}'. Supported: {', '.join(config.SUPPORTED_EXTENSIONS)}"
    )


def extract_file(path: str, page_markers: bool = True, clean: bool = True) -> str:
    """Text of a knowledge-base document, pre-processed (see src/preprocess.py) unless clean=False."""
    with open(path, "rb") as f:
        text = extract_text(os.path.basename(path), f.read(), page_markers=page_markers)
    if clean:
        from src.preprocess import clean_text
        text = clean_text(text)
    return text
